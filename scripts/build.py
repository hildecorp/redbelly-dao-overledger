#!/usr/bin/env python3
"""Builds every output from content/report.json, content/facts.json, content/sources.json and the diagrams:
  public/index.html          home page: the short answer, key findings, how to read the labels and the list of chapters
  public/chapters/*.html     one page per chapter, served at /chapters/<name> (Vercel cleanUrls)
  public/sources.html        the source list, served at /sources: every source, the exact words relied on, and where each is cited
  public/404.html            the page shown for an unknown address
  public/site.css, site.js   shared style and behaviour (copied from scripts/)
  public/downloads/          the report as PDF (Playwright print of build/print.html) and Word (python-docx),
                             plus SOURCES.md and COMMUNITY-SUMMARY.md
  SOURCES.md, COMMUNITY-SUMMARY.md at the repository root
Diagrams are hand-written SVG in public/diagrams/ID-wide.svg and ID-narrow.svg, specified in content/diagrams.json.
"""
import html, json, re, shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "public"
SITE = "https://overledger.vercel.app"
PDF_NAME = "redbelly-overledger-feasibility-report.pdf"
DOCX_NAME = "redbelly-overledger-feasibility-report.docx"
VERIFIED = "15 Oct 2026"

R = json.loads((ROOT / "content/report.json").read_text())
FACTS = {f["id"]: f for f in json.loads((ROOT / "content/facts.json").read_text())}
SOURCES = {s["id"]: s for s in json.loads((ROOT / "content/sources.json").read_text())}
DIAGRAMS = {d["id"]: d for d in json.loads((ROOT / "content/diagrams.json").read_text())}
CITE = re.compile(r"\{([a-z0-9 -]+)\}")
LABELS = {"confirmed": "Confirmed", "tested": "Tested", "analysis": "Analysis", "unknown": "Unknown"}
# Where each diagram sits: (section id, index of the block it follows; -1 = straight after the section intro)
FIG_AT = {"stack": ("overledger", 0), "routes": ("how-connect", 0), "access": ("redbelly", 2), "levels": ("now", -1),
          "work": ("work", -1), "receptor": ("compliance", 0), "breaks": ("compliance", 1)}
GROUPS = [("g-brief", "The task brief", ["brief"]),
          ("g-tch", "The Clearing House and Quant", ["tch", "quant-pr"]),
          ("g-ol", "Overledger documentation", ["ol-"]),
          ("g-rb", "Redbelly technical documentation", ["rb-"]),
          ("g-vine", "Vine, Redbelly developer portal", ["vine-"]),
          ("g-bridges", "Bridge announcements", ["celer", "x-wrbnt"]),
          ("g-tk", "Tokeniser", ["tk-"]),
          ("g-task", "Earlier Task Board research", ["task"]),
          ("g-test", "Our own tests", ["test-"])]

def group_of(sid):
    for gid, _, keys in GROUPS:
        if any(sid == k or (k.endswith("-") and sid.startswith(k)) or (k == "task" and sid.startswith("task")) for k in keys):
            return gid
    raise SystemExit(f"source {sid} has no group")

ICON = {
 "check": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M5 12.5l4.5 4.5L19 7.5"/></svg>',
 "tested": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M9 3h6M10 3v6l-5 9a2 2 0 0 0 1.8 3h10.4a2 2 0 0 0 1.8-3l-5-9V3"/></svg>',
 "analysis": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M4 19l5-6 4 3 7-9"/></svg>',
 "unknown": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><circle cx="12" cy="12" r="9"/><path d="M9.5 9.5a2.5 2.5 0 1 1 3.5 2.3c-.6.3-1 .9-1 1.6v.6M12 17.2v.1"/></svg>',
 "notyet": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" aria-hidden="true"><path d="M6 6l12 12M18 6L6 18"/></svg>',
 "none": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" aria-hidden="true"><path d="M6 12h12"/></svg>',
 "info": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" aria-hidden="true"><circle cx="12" cy="12" r="9.5"/><path d="M12 11v6M12 7.5v.1"/></svg>',
 "pdf": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M12 4v11M7 10l5 5 5-5M5 20h14"/></svg>',
 "doc": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M14 3H7a2 2 0 0 0-2 2v14a2 2 0 0 0 2 2h10a2 2 0 0 0 2-2V8z"/><path d="M14 3v5h5M9 13h6M9 17h6"/></svg>',
 "list": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" aria-hidden="true"><path d="M9 6h11M9 12h11M9 18h11M4.5 6h.1M4.5 12h.1M4.5 18h.1"/></svg>',
 "chat": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M21 12a8 8 0 0 1-11.6 7.1L4 20l1-4.6A8 8 0 1 1 21 12z"/></svg>',
 "right": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M5 12h14M13 6l6 6-6 6"/></svg>',
 "left": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M19 12H5M11 6l-6 6 6 6"/></svg>',
 "up": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M12 19V5M6 11l6-6 6 6"/></svg>',
 "copy": '<svg class="ic-copy" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><rect x="9" y="9" width="11" height="11" rx="2"/><path d="M5 15V6a2 2 0 0 1 2-2h8"/></svg><svg class="ic-done" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M5 12.5l4.5 4.5L19 7.5"/></svg>',
 "expand": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M15 3h6v6M9 21H3v-6M21 3l-7 7M3 21l7-7"/></svg>',
 "home": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M4 11l8-7 8 7M6 9.5V20h12V9.5"/></svg>',
 "chev": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M6 9l6 6 6-6"/></svg>',
 "close": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2" stroke-linecap="round" aria-hidden="true"><path d="M6 6l12 12M18 6L6 18"/></svg>',
 "search": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" aria-hidden="true"><circle cx="11" cy="11" r="7"/><path d="M20 20l-3.5-3.5"/></svg>',
 "repo": '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M4 19.5A2.5 2.5 0 0 1 6.5 17H20V3H6.5A2.5 2.5 0 0 0 4 5.5v14z"/><path d="M8 7h8"/></svg>',
}

def esc(s):
    return html.escape(s, quote=True)

def strip_num(title):
    return re.sub(r"^\d+\.\s*", "", title)

def nn(i):
    return f"{i:02d}"

# ---------- walk the report in reading order: number sources, figures and record where each fact is used ----------
ORDER, USED_IN, FIGNUM = [], {}, {}
SEC = {s["id"]: s for s in R["sections"]}
SEC_NUM = {s["id"]: i + 1 for i, s in enumerate(R["sections"])}

def figs_for(sec_id, after):
    return [d for d, (s, a) in FIG_AT.items() if s == sec_id and a == after and d in DIAGRAMS]

def note(text, where):
    for m in CITE.finditer(text):
        for fid in m.group(1).split():
            if fid not in FACTS:
                raise SystemExit(f"unknown fact {fid} in {where}")
            sid = FACTS[fid]["source"]
            if sid not in ORDER:
                ORDER.append(sid)
            USED_IN.setdefault(fid, [])
            if where not in USED_IN[fid]:
                USED_IN[fid].append(where)

def walk():
    note(R["status_note"], "top")
    for k in R["key_findings"]:
        note("{" + k["cite"] + "}", "findings")
    for s in R["sections"]:
        note(s.get("intro", ""), s["id"])
        def fig(after):
            for d in figs_for(s["id"], after):
                FIGNUM[d] = len(FIGNUM) + 1
                note(DIAGRAMS[d]["caption"], s["id"])
        fig(-1)
        for i, b in enumerate(s["blocks"]):
            note(b.get("text", ""), s["id"])
            for li in b.get("list", []):
                note(li, s["id"])
            for row in b.get("rows", []):
                note("{" + row[2] + "}", s["id"])
            fig(i)
walk()
NUM = {sid: i + 1 for i, sid in enumerate(ORDER)}
USED = set(USED_IN)

def where_label(w):
    if w == "top":
        return ("Opening note", "top")
    if w == "findings":
        return ("Key findings", "findings")
    return (f"{SEC_NUM[w]}. {SEC[w]['short']}", w)

# ---------- citations ----------
def cite_html(ids, mode):
    by = {}
    for fid in ids.split():
        by.setdefault(NUM[FACTS[fid]["source"]], []).append(fid)
    out = []
    for n in sorted(by):
        href = f"#s{n}" if mode == "print" else f"/sources#s{n}"
        out.append(f'<a href="{href}" data-n="{n}" data-f="{" ".join(by[n])}" aria-label="Source {n}">S{n}</a>')
    return '<sup class="cite">' + "".join(out) + "</sup>"

def rich(t, mode="web"):
    out, last = [], 0
    for m in CITE.finditer(t):
        before = esc(t[last:m.start()]).rstrip()
        cut = before.rfind(" ") + 1
        out.append(before[:cut] + '<span class="nw">' + before[cut:] + cite_html(m.group(1), mode) + "</span>")
        last = m.end()
    out.append(esc(t[last:]))
    return "".join(out)

def plain(t):
    def rep(m):
        nums = sorted({NUM[FACTS[f]["source"]] for f in m.group(1).split()})
        return " [" + ", ".join(f"S{n}" for n in nums) + "]"
    return re.sub(r"\s*\{([a-z0-9 -]+)\}", rep, t)

def url_html(u):
    """Visible URL: scheme kept in one unbreakable piece, break hints only after path slashes, so the PDF text layer reads https://host/path."""
    scheme, sep, rest = u.partition("://")
    if not sep:
        return esc(u)
    parts = [esc(x) for x in rest.split("/")]
    return '<span class="sc">' + esc(scheme) + '://</span>' + "/<wbr>".join(parts)

def src_url(s):
    return f"{SITE}/{s['url']}" if s.get("local") else s["url"]

def pill(kind, text=None):
    text = text or LABELS.get(kind, kind)
    return f'<span class="pill {kind}"><span class="dot" aria-hidden="true"></span>{esc(text)}</span>'

def svg_inline(did, variant):
    p = PUB / "diagrams" / f"{did}-{variant}.svg"
    if not p.exists():
        return f'<svg viewBox="0 0 960 120" role="img"><title>{esc(DIAGRAMS[did]["title"])}</title><rect x="1" y="1" width="958" height="118" rx="8" fill="#17222a" stroke="#263641"/><text x="480" y="66" text-anchor="middle" font-family="Be Vietnam Pro, sans-serif" font-size="18" fill="#8fa1ad">Diagram in progress</text></svg>'
    s = p.read_text()
    return re.sub(r"<\?xml[^>]*>\s*", "", s).strip()

def figure_html(did, mode):
    d = DIAGRAMS[did]; n = FIGNUM[did]
    zoom = "" if mode == "print" else f'<button class="zoom" type="button" aria-label="Open figure {n} full screen">{ICON["expand"]}Full screen</button>'
    narrow = "" if mode == "print" else f'<div class="art narrow">{svg_inline(did, "narrow")}</div>'
    return (f'<figure class="fig rv" id="fig-{did}"><div class="fhead"><div><div class="fk">Figure {n}</div><p class="ft">{esc(d["title"])}</p></div>{zoom}</div>'
            f'<div class="art wide">{svg_inline(did, "wide")}</div>{narrow}'
            f'<figcaption>{rich(d["caption"], mode)}</figcaption></figure>')

def table_row(r, mode):
    ans = r[1]
    if ans.startswith("Not possible yet"):
        kind, head = "notyet", "Not possible yet"
    elif ans.startswith("Unknown"):
        kind, head = "unknown", "Unknown"
    else:
        kind, head = "none", ans.split(",")[0]
    rest = ans[len(head):].lstrip(", ")
    why = (rest[0].upper() + rest[1:]) if rest else ""
    return (f'<tr><td>{esc(r[0])}</td><td>{pill(kind, head)}<span class="why">{esc(why)}{cite_html(r[2], mode)}</span></td></tr>')

def block_html(s, b, mode):
    lab = b["label"]
    hx = "h2" if mode == "web" else "h3"
    if lab == "summary_table":
        return ('<table class="tbl rv"><thead><tr><th scope="col">Question</th><th scope="col">Answer today</th></tr></thead><tbody>' +
                "".join(table_row(r, mode) for r in b["rows"]) + "</tbody></table>")
    cls = "card rv" + (" method" if lab == "method" else "")
    head = pill(lab) if lab in LABELS else pill("none", "How it was done")
    h = f'<div class="{cls}"><div class="chead">{head}'
    if b.get("heading"):
        h += f'<{hx}>{esc(b["heading"])}</{hx}>'
    h += "</div>"
    if b.get("text"):
        h += f'<p>{rich(b["text"], mode)}</p>'
    if b.get("list"):
        items = b["list"]
        if s["id"] == "work":
            h += '<ol class="steps">'
            for li in items:
                m = re.match(r"Level (\d), step (\d): (.*)", li)
                lvl, step, body = m.group(1), m.group(2), m.group(3)
                body = body[0].upper() + body[1:]
                h += f'<li class="l{lvl}"><span class="sn">{step}</span><span class="lvl">Level {lvl}, step {step}</span>{rich(body, mode)}</li>'
            h += "</ol>"
        elif s["id"] == "compliance" and b.get("heading", "").startswith("Where compliance"):
            h += '<ol class="breaks">' + "".join(f"<li>{rich(li, mode)}</li>" for li in items) + "</ol>"
        else:
            h += "<ul>" + "".join(f"<li>{rich(li, mode)}</li>" for li in items) + "</ul>"
    return h + "</div>"

# ---------- chapters: the website shows the key findings on the home page and each chapter on its own page ----------
def slug(t):
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")

CHAPTERS = [(s["id"], s["short"], s["kicker"]) for s in R["sections"]] + [("summary", "Community summary", "Share it")]
CH_IDS = [c[0] for c in CHAPTERS]
CH_TITLE = {c[0]: c[1] for c in CHAPTERS}
CH_SLUG = {cid: slug(t) for cid, t, _ in CHAPTERS}

def chapter_url(cid):
    return f"/chapters/{CH_SLUG[cid]}"

def chapter_top(cid):
    """Bar above each chapter: back to all chapters, "Chapter 3 of 11", and a slider that opens any chapter.
    Without JavaScript the slider is hidden and a plain progress bar shows the position."""
    i = CH_IDS.index(cid) + 1; n = len(CHAPTERS)
    urls = "|".join(chapter_url(c) for c in CH_IDS)
    titles = "|".join(esc(CH_TITLE[c]) for c in CH_IDS)
    ticks = "".join(f'<span style="left:{k / (n - 1) * 100:.2f}%"></span>' for k in range(n))
    return (f'<div class="chap-top"><a class="chap-all" href="/#chapters">{ICON["left"]}All chapters</a>'
            f'<span class="chap-of" aria-hidden="true">Chapter <b>{i}</b> of {n}</span>'
            f'<span class="chap-meter" aria-hidden="true"><span style="width:{i / n * 100:.1f}%"></span></span>'
            f'<span class="chap-slide" style="--p:{(i - 1) / (n - 1) * 100:.2f}%"><span class="chap-ticks" aria-hidden="true">{ticks}</span>'
            f'<input class="chap-range" type="range" min="1" max="{n}" step="1" value="{i}" aria-label="Chapter {i} of {n}. Slide to open another chapter"'
            f' aria-valuetext="Chapter {i} of {n}: {esc(CH_TITLE[cid])}" data-urls="{urls}" data-titles="{titles}"></span>'
            f'<span class="chap-tip" aria-live="polite" hidden></span></div>')

def neighbours(cid):
    """(href, label) for the previous and next stop: the home page before chapter 1, the source list after the last."""
    i = CH_IDS.index(cid)
    prev = (chapter_url(CH_IDS[i - 1]), CH_TITLE[CH_IDS[i - 1]]) if i > 0 else ("/#chapters", "Key findings and chapters")
    nxt = (chapter_url(CH_IDS[i + 1]), CH_TITLE[CH_IDS[i + 1]]) if i + 1 < len(CH_IDS) else ("/sources", "Source list")
    return prev, nxt

def chapter_nav(cid):
    (ph, pt), (nh, nt) = neighbours(cid)
    first, last = CH_IDS.index(cid) == 0, CH_IDS.index(cid) == len(CH_IDS) - 1
    return (f'<nav class="chap-nav" aria-label="Chapters">'
            f'<a class="chap-prev" href="{ph}" rel="prev"><span class="dir">{ICON["left"]}{"Back to" if first else "Previous"}</span><span class="t">{esc(pt)}</span></a>'
            f'<a class="chap-next" href="{nh}" rel="next"><span class="dir">{"Then" if last else "Next"}{ICON["right"]}</span><span class="t">{esc(nt)}</span></a></nav>')

def words_in(html):
    return len(re.sub(r"<[^>]+>", " ", html).split())

def chapters_html(bodies):
    cards = ""
    for i, (cid, title, kicker) in enumerate(CHAPTERS, 1):
        mins = max(1, round(words_in(bodies[cid]) / 220))
        cards += (f'<li><a class="chap-card rv" href="{chapter_url(cid)}"><span class="cn">{nn(i)}</span><span class="ck">{esc(kicker)}</span>'
                  f'<span class="ct">{esc(title)}</span><span class="cm">{mins} min read{ICON["right"]}</span></a></li>')
    return (f'<section class="s chapters" id="chapters" aria-labelledby="h-chapters"><div class="rv"><p class="kicker"><span class="bar"></span>Read it in parts</p>'
            f'<h2 id="h-chapters">Chapters</h2><p class="intro">The full report in {len(CHAPTERS)} short chapters, each on its own page. Open one, then use Next and Previous at the end of each.</p></div>'
            f'<ol class="chap-grid">{cards}</ol></section>')

def section_html(s, mode):
    i = SEC_NUM[s["id"]]
    h = (f'<section class="s" id="{s["id"]}" aria-labelledby="h-{s["id"]}">' + (chapter_top(s["id"]) if mode == "web" else "") + '<div class="rv"><p class="kicker"><span class="bar"></span>'
         f'<span class="num">{nn(i)}</span>{esc(s["kicker"])}</p>{"<h1" if mode == "web" else "<h2"} id="h-{s["id"]}">{esc(strip_num(s["title"]))}{"</h1>" if mode == "web" else "</h2>"}')
    if s.get("intro"):
        h += f'<p class="intro">{rich(s["intro"], mode)}</p>'
    h += "</div>"
    for d in figs_for(s["id"], -1):
        h += figure_html(d, mode)
    for bi, b in enumerate(s["blocks"]):
        h += block_html(s, b, mode)
        for d in figs_for(s["id"], bi):
            h += figure_html(d, mode)
    return h + (chapter_nav(s["id"]) if mode == "web" else "") + "</section>"

def findings_html(mode):
    h = '<section class="s" id="findings" aria-labelledby="h-findings"><div class="rv"><p class="kicker"><span class="bar"></span><span class="num">00</span>Short answer</p><h2 id="h-findings">Key findings</h2></div><div class="verdicts">'
    for i, k in enumerate(R["key_findings"], 1):
        cls = "none" if k["verdict"] == "None" else k["label"]
        h += (f'<article class="verdict rv {cls}"><div class="vtop"><span class="vnum">{nn(i)}</span>{pill(k["label"])}</div>'
              f'<p class="q">{esc(k["q"])}</p><p class="big">{esc(k["verdict"])}</p><p class="a">{rich(k["a"] + " {" + k["cite"] + "}", mode)}</p></article>')
    return h + "</div></section>"

def glossary_html(mode):
    items = "".join(f'<div class="gl-i"><dt>{esc(term)}</dt><dd>{rich(d, mode)}</dd></div>' for term, d in R["glossary"])
    hx = "h2" if mode == "web" else "h3"
    return (f'<div class="card rv glossary" id="glossary"><div class="chead"><{hx}>Plain words</{hx}></div>'
            f'<p class="gl-note">The technical terms this report uses, in one or two lines each. Where a term is Overledger\'s or Redbelly\'s own, the source is cited.</p><dl>{items}</dl></div>')

def post_count(t):
    return len(t.replace("{site}", "x" * 23))

def summary_html(mode):
    cs = R["community_summary"]; n = len(R["sections"]) + 1
    paras = "".join(f'<p class="sum-p" data-raw="{esc(p.replace("{site}", SITE))}">{esc(p.replace("{site}", SITE))}</p>' for p in cs["paragraphs"])
    tag = "h1" if mode == "web" else "h2"
    h = (f'<section class="s" id="summary" aria-labelledby="h-summary">' + (chapter_top("summary") if mode == "web" else "") + f'<div class="rv"><p class="kicker"><span class="bar"></span><span class="num">{nn(n)}</span>Share it</p>'
         f'<{tag} id="h-summary">Community summary</{tag}><p class="intro">Plain language, ready for the DAO to share.{"" if mode == "print" else " Copy the paragraphs, a single post, or the whole thread."}</p></div>')
    btn = lambda attrs, label, msg, extra="": (f'<button class="copy{extra}" type="button" {attrs} data-msg="{msg}">{ICON["copy"]}<span class="lbl">{label}</span></button>')
    hx = "h2" if mode == "web" else "h3"
    h += (f'<div class="share rv"><div class="shead"><{hx}>A few paragraphs</{hx}>' +
          ("" if mode == "print" else btn('data-copy=".sum-p"', "Copy summary", "Summary copied", " primary")) + f'</div><div class="body">{paras}</div></div>')
    h += (f'<div class="share rv"><div class="shead"><{hx}>As an X thread, {len(cs["thread"])} posts</{hx}>' +
          ("" if mode == "print" else btn('data-copy=".post-text"', "Copy all posts", "Thread copied", " primary")) + '</div><ol class="posts">')
    for i, t in enumerate(cs["thread"], 1):
        txt = t.replace("{site}", SITE); c = post_count(t)
        h += (f'<li class="post"><div class="ptop"><span>Post {i} of {len(cs["thread"])}</span><span class="count{" full" if c > 260 else ""}">{c} / 280</span></div>'
              f'<p class="post-text" data-raw="{esc(txt)}">{esc(txt)}</p>' +
              ("" if mode == "print" else '<div class="pbot">' + btn('data-text="' + esc(txt) + '"', "Copy post", f"Post {i} copied") + "</div>") + "</li>")
    return h + "</ol></div>" + (chapter_nav("summary") if mode == "web" else "") + "</section>"

# ---------- shared page parts ----------
CSS_V = "?v=" + __import__("hashlib").sha256((ROOT / "scripts/site.css").read_bytes() + (ROOT / "scripts/site.js").read_bytes()).hexdigest()[:10]

def head(title, desc, path, extra=""):
    url = SITE + path
    return (f'<!doctype html><html lang="en-AU"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{esc(title)}</title><meta name="description" content="{esc(desc)}"><meta name="theme-color" content="#0c1317">'
            f'<link rel="canonical" href="{url}"><meta property="og:title" content="{esc(title)}"><meta property="og:description" content="{esc(desc)}">'
            f'<meta property="og:type" content="article"><meta property="og:url" content="{url}"><meta property="og:site_name" content="Redbelly DAO community research">'
            f'<meta property="og:image" content="{SITE}/brand/og.png"><meta property="og:image:width" content="1200"><meta property="og:image:height" content="630">'
            f'<meta property="og:image:alt" content="{esc(R["title"])} Feasibility research, sources last verified {VERIFIED}.">'
            f'<meta name="twitter:card" content="summary_large_image">'
            f'<link rel="icon" type="image/png" sizes="32x32" href="/brand/icon-32.png"><link rel="icon" type="image/png" sizes="512x512" href="/brand/icon.png">'
            f'<link rel="apple-touch-icon" href="/brand/icon-180.png">'
            f'<link rel="preload" href="/fonts/be-vietnam-pro-latin-400-normal.woff2" as="font" type="font/woff2" crossorigin>'
            f'<link rel="stylesheet" href="/site.css{CSS_V}">{extra}</head>')

def header(current):
    def link(href, label, key):
        cur = ' aria-current="page"' if key == current else ""
        return f'<a href="{href}"{cur}>{label}</a>'
    return ('<a class="skip" href="#main">Skip to content</a><div class="progress" aria-hidden="true"><span></span></div>'
            '<header class="top"><a class="brand" href="/"><img src="/brand/dao-logo.png" alt="Redbelly DAO, home" width="62" height="44"></a>'
            f'<nav aria-label="Pages">{link("/", "Report", "report")}{link("/sources", "Sources", "sources")}{link("/evidence", "Evidence", "evidence")}</nav></header>')

def toc_html(items, extra_links, head_text):
    lis = "".join(f'<li><a class="rn-item" href="#{i}" data-t="{i}"><span class="rn-l">{n}</span><span class="rn-t">{esc(t)}</span></a></li>' for i, n, t in items)
    chips = "".join(f'<li><a href="#{i}" data-t="{i}"><span class="n">{n}</span>{esc(t)}</a></li>' for i, n, t in items)
    return (f'<aside class="toc rnav" aria-label="{esc(head_text)}"><p class="toc-head">{esc(head_text)}</p><ol class="rn">{lis}</ol>'
            f'<div class="toc-links">{extra_links}</div></aside>', f'<nav class="chips" aria-label="{esc(head_text)}"><ol>{chips}</ol></nav>')

def foot():
    return ('<footer class="foot"><div><p>Feasibility research only. Redbelly and Quant have no announced partnership or talks, and nothing on this site suggests otherwise.</p>'
            f'<p>Sources last verified {VERIFIED}. Every claim links to the source list. Saved copies of every source and our test logs are under <a href="/evidence">/evidence</a>.</p></div>'
            '<div><p>Fonts: Be Vietnam Pro and JetBrains Mono, SIL Open Font License 1.1. Logo and icon: Redbelly DAO, from the Task Board brand kit.</p>'
            '<p>Diagrams drawn for this report. No third-party images.</p></div></footer>')

DOWNLOADS = (f'<a href="/downloads/{PDF_NAME}">{ICON["pdf"]}Download PDF</a><a href="/downloads/{DOCX_NAME}">{ICON["doc"]}Word file</a>'
             f'<a href="/sources">{ICON["list"]}Source list</a>')

def report_nav(active):
    """Side menu: two entries, Key findings and Chapters, with the chapters as a sub list.
    active is 'home' on the home page, or a chapter id on a chapter page."""
    on_home = active == "home"
    subs = ""
    for i, (cid, t, _) in enumerate(CHAPTERS, 1):
        cur = ' class="active" aria-current="page"' if cid == active else ""
        subs += f'<li><a href="{chapter_url(cid)}"{cur}><span class="n">{nn(i)}</span><span>{esc(t)}</span></a></li>'
    fhref, chref = ("#findings", "#chapters") if on_home else ("/#findings", "/#chapters")
    open_cls = "" if on_home else " open"
    return (f'<aside class="toc rnav" aria-label="Report contents"><p class="toc-head">Report</p><ol class="rn">'
            f'<li><a class="rn-item" href="{fhref}" data-t="findings"><span class="rn-l">A</span><span class="rn-t">Key findings</span></a></li>'
            f'<li class="rn-group{open_cls}"><div class="rn-row"><a class="rn-item{"" if on_home else " here"}" href="{chref}" data-t="chapters"><span class="rn-l">B</span><span class="rn-t">Chapters</span>'
            f'<span class="cnt">{len(CHAPTERS)}</span></a><button class="rn-tog" type="button" aria-expanded="{"false" if on_home else "true"}" aria-controls="rn-sub" aria-label="Show or hide the chapters">{ICON["chev"]}</button></div>'
            f'<div class="rn-subw" id="rn-sub"><ol class="rn-sub">{subs}</ol></div></li></ol>'
            f'<div class="toc-links">{DOWNLOADS}</div></aside>')

def chapter_sheet(active):
    items = "".join(f'<li><a href="{chapter_url(cid)}"{" class=\"active\" aria-current=\"page\"" if cid == active else ""}><span class="n">{nn(i)}</span>{esc(t)}</a></li>'
                    for i, (cid, t, _) in enumerate(CHAPTERS, 1))
    return f'<div class="sheet" id="sheet" hidden><p class="sh-head">Chapters</p><ol>{items}</ol></div>'

def dock(active):
    """Bottom bar. Home: Findings, Chapters, Sources, Top. Chapter: Chapters, previous, position, next, Sources, Top."""
    chap = f'<button class="dk dk-chap" type="button" aria-expanded="false" aria-controls="sheet">{ICON["list"]}<span class="dl">Chapters</span></button>'
    src = f'<span class="dk-sep" aria-hidden="true"></span><a class="dk dk-src" href="/sources">{ICON["repo"]}<span class="dl">Sources</span></a>'
    top = f'<button class="dk dk-top" type="button" aria-label="Back to top">{ICON["up"]}</button>'
    if active == "home":
        mid = f'<a class="dk dk-home" href="#findings">{ICON["home"]}<span class="dl">Findings</span></a>' + chap
    else:
        (ph, _), (nh, nt) = neighbours(active)
        i = CH_IDS.index(active) + 1
        mid = (chap + f'<span class="dk-mid"><a class="dk dk-prev" href="{ph}" aria-label="Previous">{ICON["left"]}</a>'
               f'<span class="dk-pos">{i} / {len(CHAPTERS)}</span><a class="dk dk-next" href="{nh}" aria-label="Next: {esc(nt)}">{ICON["right"]}</a></span>')
    return (f'<nav class="dock" aria-label="Quick navigation">{mid}{src}{top}</nav>' + chapter_sheet(active) +
            '<div class="toast" role="status" aria-live="polite"></div>')

def dock_sources():
    return ('<nav class="dock" aria-label="Quick navigation">'
            f'<a class="dk" href="/">{ICON["left"]}<span class="dl">Back to the report</span></a><span class="dk-sep" aria-hidden="true"></span>'
            f'<a class="dk dk-find" href="#q">{ICON["search"]}<span class="dl">Filter</span></a>'
            f'<button class="dk dk-top" type="button" aria-label="Back to top">{ICON["up"]}</button></nav>'
            '<div class="toast" role="status" aria-live="polite"></div>')

HERO_ART = ('<svg class="hero-art" viewBox="0 0 300 220" aria-hidden="true" focusable="false">'
            '<g stroke="#345062" stroke-width="1.2" fill="none">'
            '<path d="M40 60L120 40L200 70L260 40M40 60L90 130L170 120L200 70M90 130L60 190M170 120L240 160L260 40M170 120L150 190"/></g>'
            '<g fill="#17222a" stroke="#4a6375" stroke-width="1.5"><circle cx="40" cy="60" r="7"/><circle cx="120" cy="40" r="6"/><circle cx="260" cy="40" r="6"/>'
            '<circle cx="90" cy="130" r="6"/><circle cx="60" cy="190" r="5"/><circle cx="240" cy="160" r="7"/><circle cx="150" cy="190" r="5"/></g>'
            '<circle cx="200" cy="70" r="9" fill="#1d3140" stroke="#EF5350" stroke-width="2"/><circle cx="170" cy="120" r="9" fill="#1d3140" stroke="#8fa1ad" stroke-width="1.5" stroke-dasharray="3 3"/>'
            '</svg>')

def cite_data(page_html):
    """Source titles and quoted words for the citation previews, limited to what the page cites."""
    nums = sorted({int(n) for n in re.findall(r'data-n="(\d+)"', page_html)})
    fids = sorted({f for grp in re.findall(r'data-f="([^"]+)"', page_html) for f in grp.split()})
    by_num = {NUM[s]: s for s in ORDER}
    return json.dumps({"sources": {str(n): {"p": SOURCES[by_num[n]]["publisher"], "t": SOURCES[by_num[n]]["title"]} for n in nums},
                       "facts": {f: FACTS[f]["quote"] for f in fids}}, ensure_ascii=False).replace("</", "<\\/")

VIEWER = ('<div class="viewer" aria-hidden="true" role="dialog" aria-label="Figure, full screen"><button class="close" type="button" aria-label="Close">'
          + ICON["close"] + '</button><div class="vin"></div></div>')

def page(title, desc, path, current, inner, tail, extra_head=""):
    js = f'<script src="/site.js{CSS_V}"></script>'
    data = f'<script type="application/json" id="cite-data">{cite_data(inner)}</script>' if 'class="cite"' in inner else ""
    viewer = VIEWER if 'class="zoom"' in inner else ""
    return (head(title, desc, path, extra_head) + "<body>" + header(current) + f'<main id="main" tabindex="-1">{inner}</main>' +
            viewer + tail + data + js + "</body></html>")

def chapter_bodies():
    b = {s["id"]: section_html(s, "web") for s in R["sections"]}
    b["summary"] = summary_html("web")
    return b

def home_page(bodies):
    nsrc, nfacts = len(ORDER), len(USED)
    hero = (f'<section class="hero" aria-labelledby="title">{HERO_ART}<p class="kicker"><span class="bar"></span>Redbelly DAO community research</p>'
            f'<h1 id="title">{esc(R["title"])}</h1><p class="lead">{esc(R["subtitle"])}. What Overledger is, whether Redbelly can reach it, what that would take, and where compliance could break.</p>'
            f'<div class="notice">{ICON["info"]}<p>{rich(R["status_note"])}</p></div>'
            f'<p class="meta-row"><span>Last verified <b>{VERIFIED}</b></span><span><b>{nsrc}</b> sources</span><span><b>{nfacts}</b> quoted facts</span><span><b>{len(FIGNUM)}</b> diagrams</span><span><b>{len(CHAPTERS)}</b> chapters</span></p>'
            f'<div class="actions"><a class="btn primary" href="{chapter_url(CH_IDS[0])}">{ICON["right"]}Start reading</a>'
            f'<a class="btn" href="/downloads/{PDF_NAME}">{ICON["pdf"]}Download PDF</a><a class="btn" href="/downloads/{DOCX_NAME}">{ICON["doc"]}Word file</a>'
            f'<a class="btn" href="{chapter_url("summary")}">{ICON["chat"]}Community summary</a><a class="btn" href="/sources">{ICON["list"]}Source list</a></div></section>')
    legend = ('<div class="card rv legend-card"><div class="chead"><h3>How to read the labels</h3></div><ul class="legend">' +
              "".join(f'<li>{pill(k)}<span>{esc(v)}</span></li>' for k, v in R["labels"].items()) + "</ul></div>")
    inner = (f'{hero}<div class="layout">{report_nav("home")}<div class="content">' +
             findings_html("web") + legend + glossary_html("web") + chapters_html(bodies) + foot() + "</div></div>")
    first = chapter_url(CH_IDS[0])
    return page(R["title"] + " | Redbelly DAO research", R["subtitle"], "/", "report", inner, dock("home"),
                f'<link rel="prefetch" href="{first}">')

def chapter_page(cid, body):
    i = CH_IDS.index(cid) + 1
    (ph, _), (nh, _) = neighbours(cid)
    title = f"{i}. {CH_TITLE[cid]} | {R['title']}"
    s = SEC.get(cid)
    desc = s["description"] if s else R["community_summary"]["description"]
    inner = (f'<div class="layout">{report_nav(cid)}<div class="content chapter">{body}{foot()}</div></div>')
    extra = f'<link rel="prev" href="{ph}"><link rel="next" href="{nh}"><link rel="prefetch" href="{nh}">'
    return page(title, desc, chapter_url(cid), "report", inner, dock(cid), extra)

def not_found_page():
    inner = (f'<section class="hero" aria-labelledby="title">{HERO_ART}<p class="kicker"><span class="bar"></span>Page not found</p>'
             f'<h1 id="title">That page is not part of this report</h1><p class="lead">The address may be mistyped, or it may point to an older version of the site.</p>'
             f'<div class="actions"><a class="btn primary" href="/">{ICON["home"]}Key findings</a><a class="btn" href="/#chapters">{ICON["list"]}All chapters</a>'
             f'<a class="btn" href="/sources">{ICON["repo"]}Source list</a></div></section>')
    return page("Page not found | " + R["title"], "This page is not part of the report.", "/404", "", inner, "")

def snapshot_info(s):
    import hashlib
    p = (ROOT / "evidence/snapshots" / s["file"]).resolve()
    rel = p.relative_to(ROOT)
    man = {m["file"]: m for m in json.loads((ROOT / "evidence/manifest.json").read_text())}
    m = man.get(str(rel))
    sha = m["sha256"] if m else hashlib.sha256(p.read_bytes()).hexdigest()
    when = (m["captured_utc"].replace("T", " ").replace("Z", " UTC")) if m else ""
    return str(rel), sha, when

def used_href(w):
    if w == "top":
        return "/"
    if w == "findings":
        return "/#findings"
    return chapter_url(w)

def sources_page():
    items, groups = [], []
    for gid, gname, _ in GROUPS:
        srcs = [s for s in ORDER if group_of(s) == gid]
        srcs.sort(key=lambda s: NUM[s])
        if srcs:
            items.append((gid, nn(len(items) + 1), gname)); groups.append((gid, gname, srcs))
    toc, chips = toc_html(items, f'<a href="/">{ICON["left"]}Back to the report</a><a href="/downloads/SOURCES.md">{ICON["doc"]}Source list as Markdown</a>', "Sources")
    hero = (f'<section class="hero" aria-labelledby="title"><p class="kicker"><span class="bar"></span>Source list</p><h1 id="title">Where every fact comes from</h1>'
            f'<p class="lead">The report cites each source as S1, S2 and so on. Under each one are the exact words the report relies on, so you can find them on the page with Ctrl+F, and links back to where each is used.</p>'
            f'<div class="notice">{ICON["info"]}<p>A text copy of every page as it read on {VERIFIED} is published here under /evidence/snapshots, with its SHA-256 in <a href="/evidence/manifest.json">evidence/manifest.json</a>. The check script fails if any quote is missing from its copy.</p></div>'
            f'<p class="meta-row"><span><b>{len(ORDER)}</b> sources</span><span><b>{len(USED)}</b> quoted facts</span><span>Last verified <b>{VERIFIED}</b></span></p>'
            f'<div class="actions"><a class="btn primary" href="/">{ICON["left"]}Back to the report</a><a class="btn" href="/downloads/{PDF_NAME}">{ICON["pdf"]}Download PDF</a>'
            f'<a class="btn" href="/downloads/SOURCES.md">{ICON["doc"]}Source list as Markdown</a></div></section>')
    body = (f'<div class="filter rv">{ICON["search"]}<label class="sr" for="q">Filter sources</label><input id="q" type="search" placeholder="Filter by publisher, title or quoted words" autocomplete="off">'
            f'<span class="fc" aria-live="polite">{len(ORDER)} of {len(ORDER)} sources</span></div><p class="empty">No source matches that filter.</p>')
    for gid, gname, srcs in groups:
        body += f'<section class="group" id="{gid}"><h2>{esc(gname)}<span class="gn">{len(srcs)} {"source" if len(srcs) == 1 else "sources"}</span></h2>'
        for sid in srcs:
            s = SOURCES[sid]; n = NUM[sid]
            if s.get("local"):
                url = f'<a class="su" href="/{esc(s["url"])}">{esc(src_url(s))}</a>'
            else:
                url = f'<a class="su" href="{esc(s["url"])}" target="_blank" rel="noopener">{esc(s["url"])}</a>'
            rel, sha, when = snapshot_info(s)
            body += (f'<article class="src" id="s{n}"><div class="sh"><span class="sn">S{n}</span><div><span class="sp">{esc(s["publisher"])}</span>'
                     f'<p class="st">{esc(s["title"])}</p>{url}</div></div>'
                     f'<p class="snap">Saved copy: <a href="/{esc(rel)}">{esc(rel)}</a> · captured {esc(when)} · SHA-256 {sha[:12]}…{sha[-6:]}</p><ul class="qs">')
            for f in (f for f in FACTS.values() if f["source"] == sid and f["id"] in USED):
                used = ", ".join(f'<a href="{used_href(w)}">{esc(where_label(w)[0])}</a>' for w in USED_IN[f["id"]])
                body += (f'<li><div><q>{esc(f["quote"])}</q><span class="used">Used in {used}</span></div>'
                         f'<button class="copy" type="button" data-text="{esc(f["quote"])}" data-msg="Quote copied" aria-label="Copy this quote">{ICON["copy"]}<span class="lbl">Copy</span></button></li>')
            body += "</ul></article>"
        body += "</section>"
    inner = f'{hero}<div class="layout">{toc}<div class="content">{chips}{body}{foot()}</div></div>'
    return page("Source list | " + R["title"], "Every source behind the Overledger and Redbelly feasibility report, with the exact words relied on.",
                "/sources", "sources", inner, dock_sources())

# ---------- print version for the PDF ----------
def drop_small_screen_rules(css):
    """An A4 page is narrower than the phone breakpoint, so the PDF must not pick up the small-screen layout."""
    out, i = [], 0
    while True:
        j = css.find("@media (max-width", i)
        if j < 0:
            out.append(css[i:]); break
        out.append(css[i:j])
        k = css.index("{", j); depth = 1; k += 1
        while depth:
            depth += {"{": 1, "}": -1}.get(css[k], 0); k += 1
        i = k
    return "".join(out)

def print_page():
    fonts = (PUB / "fonts").as_uri()
    css = drop_small_screen_rules((ROOT / "scripts/site.css").read_text()).replace('url("fonts/', f'url("{fonts}/') + (ROOT / "scripts/print.css").read_text()
    logo = (PUB / "brand/dao-logo.png").as_uri()
    hero = (f'<section class="hero"><img class="plogo" src="{logo}" alt="Redbelly DAO"><p class="kicker"><span class="bar"></span>Redbelly DAO community research · Feasibility report · {VERIFIED}</p>'
            f'<h1>{esc(R["title"])}</h1><p class="lead">{esc(R["subtitle"])}.</p><div class="notice">{ICON["info"]}<p>{rich(R["status_note"], "print")}</p></div>'
            f'<p class="meta-row"><span><b>{len(ORDER)}</b> sources</span><span><b>{len(USED)}</b> quoted facts</span><span>Web version <a href="{SITE}"><b>{SITE.replace("https://", "")}</b></a></span></p></section>')
    legend = ('<div class="card legend-card"><div class="chead"><h3>How to read the labels</h3></div><ul class="legend">' +
              "".join(f'<li>{pill(k)}<span>{esc(v)}</span></li>' for k, v in R["labels"].items()) + "</ul></div>")
    body = findings_html("print") + legend + glossary_html("print") + "".join(section_html(s, "print") for s in R["sections"]) + summary_html("print")
    src = f'<section class="s srcs" id="sources"><p class="kicker"><span class="bar"></span>Sources</p><h2>Source list</h2><p class="intro">Each source with the exact words this report relies on. Saved copies of every page, with SHA-256 hashes, are at <a href="{SITE}/evidence">{SITE.replace("https://", "")}/evidence</a>.</p><ol class="plist">'
    for sid in ORDER:
        s = SOURCES[sid]; n = NUM[sid]
        src += (f'<li id="s{n}"><b>S{n}</b> <span>{esc(s["publisher"])}: {esc(s["title"])}</span><br><a class="u" href="{esc(src_url(s))}">{url_html(src_url(s))}</a><ul>' +
                "".join(f'<li>“{esc(f["quote"])}”</li>' for f in FACTS.values() if f["source"] == sid and f["id"] in USED) + "</ul></li>")
    src += "</ol></section>"
    return (f'<!doctype html><html lang="en-AU"><head><meta charset="utf-8"><title>{esc(R["title"])}</title><style>{css}</style></head>'
            f'<body class="print"><main>{hero}{body}{src}</main></body></html>')

# ---------- Word ----------
def render_pngs():
    out = {}
    from playwright.sync_api import sync_playwright
    fonts = (PUB / "fonts").as_uri()
    face = "".join(f'@font-face{{font-family:"{fam}";src:url({fonts}/{f});font-weight:{w}}}' for fam, f, w in
                   [("Be Vietnam Pro", "be-vietnam-pro-latin-400-normal.woff2", 400), ("Be Vietnam Pro", "be-vietnam-pro-latin-500-normal.woff2", 500),
                    ("Be Vietnam Pro", "be-vietnam-pro-latin-600-normal.woff2", 600), ("Be Vietnam Pro", "be-vietnam-pro-latin-700-normal.woff2", 700),
                    ("JetBrains Mono", "jetbrains-mono-latin-400-normal.woff2", 400), ("JetBrains Mono", "jetbrains-mono-latin-500-normal.woff2", 500)])
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1000, "height": 800}, device_scale_factor=2)
        for did in FIGNUM:
            tmp = ROOT / "build" / f"fig-{did}.html"
            tmp.write_text(f'<html><head><meta charset="utf-8"><style>{face}body{{margin:0;background:#141f26}}.w{{width:960px;padding:20px}}svg{{display:block;width:100%;height:auto}}</style></head>'
                           f'<body><div class="w">{svg_inline(did, "wide")}</div></body></html>')
            pg.goto(tmp.as_uri()); pg.evaluate("document.fonts.ready.then(() => true)"); pg.wait_for_timeout(300)
            f = ROOT / "build" / f"fig-{did}.png"; pg.locator(".w").screenshot(path=str(f)); out[did] = f
        b.close()
    return out

def docx_link(par, url, text=None):
    """Adds a real, clickable hyperlink run to a Word paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "B3272A"); pr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); pr.append(u)
    r.append(pr); t = OxmlElement("w:t"); t.text = text or url; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); par._p.append(h)

def clean_app_props(path):
    """python-docx's template names a word processor in docProps/app.xml; the report was not written with it."""
    import zipfile, shutil, tempfile
    tmp = Path(tempfile.mkstemp(suffix=".docx")[1])
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/thumbnail.jpeg":
                continue
            if item.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship[^>]*thumbnail[^>]*/>', b"", data)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override[^>]*thumbnail[^>]*/>', b"", data)
            if item.filename == "docProps/app.xml":
                data = re.sub(rb"<Application>.*?</Application>", b"<Application></Application>", data)
                data = re.sub(rb"<AppVersion>.*?</AppVersion>", b"", data)
                data = re.sub(rb"<Company>.*?</Company>", b"<Company></Company>", data)
            zout.writestr(item, data)
    shutil.move(str(tmp), path)

def build_docx(path, pngs):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.4)
    for name in ("Title", "Heading 1", "Heading 2"):
        d.styles[name].font.color.rgb = RGBColor(0x0F, 0x1A, 0x22)
    st.font.size = Pt(11)
    d.add_heading(R["title"], 0)
    d.add_paragraph(R["subtitle"] + ".").runs[0].italic = True
    p0 = d.add_paragraph(f"Redbelly DAO community research. Feasibility report, sources last verified {VERIFIED}. Web version: "); docx_link(p0, SITE)
    d.add_paragraph(plain(R["status_note"]))
    d.add_heading("How to read the labels", 1)
    for k, v in R["labels"].items():
        p = d.add_paragraph(style="List Bullet"); r = p.add_run(LABELS[k] + ": "); r.bold = True; p.add_run(v)
    d.add_heading("Plain words", 1)
    d.add_paragraph("The technical terms this report uses, in one or two lines each. Where a term is Overledger's or Redbelly's own, the source is cited.")
    for term, dfn in R["glossary"]:
        p = d.add_paragraph(style="List Bullet"); r = p.add_run(term + ": "); r.bold = True; p.add_run(plain(dfn))
    d.add_heading("Key findings", 1)
    t = d.add_table(rows=1, cols=4); t.style = "Light Grid Accent 2"
    for i, x in enumerate(("Question", "Short answer", "Detail", "Status")):
        t.rows[0].cells[i].text = x
    for k in R["key_findings"]:
        c = t.add_row().cells; c[0].text = k["q"]; c[1].text = k["verdict"]; c[2].text = plain(k["a"] + " {" + k["cite"] + "}"); c[3].text = LABELS[k["label"]]
    def add_fig(did):
        d.add_picture(str(pngs[did]), width=Cm(16))
        d.inline_shapes[-1]._inline.docPr.set("descr", f"Figure {FIGNUM[did]}. {DIAGRAMS[did]['title']}. " + plain(DIAGRAMS[did]["caption"]))
        cap = d.add_paragraph(); r = cap.add_run(f"Figure {FIGNUM[did]}. {DIAGRAMS[did]['title']}. "); r.bold = True
        cap.add_run(plain(DIAGRAMS[did]["caption"]))
    for s in R["sections"]:
        d.add_heading(f"{SEC_NUM[s['id']]}. {strip_num(s['title'])}", 1)
        if s.get("intro"):
            d.add_paragraph(plain(s["intro"]))
        for did in figs_for(s["id"], -1):
            add_fig(did)
        for bi, b in enumerate(s["blocks"]):
            if b["label"] == "summary_table":
                t = d.add_table(rows=1, cols=2); t.style = "Light Grid Accent 2"
                t.rows[0].cells[0].text = "Question"; t.rows[0].cells[1].text = "Answer today"
                for r in b["rows"]:
                    c = t.add_row().cells; c[0].text = r[0]; c[1].text = plain(r[1] + " {" + r[2] + "}")
            else:
                if b.get("heading"):
                    d.add_heading(b["heading"], 2)
                tag = LABELS.get(b["label"], "Method")
                if b.get("text"):
                    p = d.add_paragraph(); r = p.add_run(f"[{tag}] "); r.bold = True; r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
                    p.add_run(plain(b["text"]))
                if b.get("list"):
                    if not b.get("text"):
                        r = d.add_paragraph().add_run(f"[{tag}]"); r.bold = True; r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
                    # numbered lists carry their own numbers, so they match the figures (Word's auto numbering would run on across lists)
                    numbered = s["id"] == "compliance" and b.get("heading", "").startswith("Where")
                    for k, li in enumerate(b["list"], 1):
                        if numbered:
                            q = d.add_paragraph(f"{k}. " + plain(li)); q.paragraph_format.left_indent = Cm(0.6); q.paragraph_format.first_line_indent = Cm(-0.6)
                        else:
                            d.add_paragraph(plain(li), style="List Bullet")
            for did in figs_for(s["id"], bi):
                add_fig(did)
    cs = R["community_summary"]
    d.add_heading(f"{len(R['sections']) + 1}. Community summary", 1)
    for p in cs["paragraphs"]:
        d.add_paragraph(p.replace("{site}", SITE))
    d.add_heading("As an X thread", 2)
    for x in cs["thread"]:
        d.add_paragraph(x.replace("{site}", SITE))
    d.add_heading("Source list", 1)
    d.add_paragraph("Every factual claim cites one of these sources. Under each source are the exact words the report relies on.")
    pr = d.add_paragraph("Saved copies of every source page, with SHA-256 hashes: "); docx_link(pr, f"{SITE}/evidence")
    for sid in ORDER:
        s = SOURCES[sid]
        p = d.add_paragraph(); r = p.add_run(f"S{NUM[sid]}  {s['publisher']}: {s['title']}"); r.bold = True
        p.paragraph_format.keep_with_next = True
        pl = d.add_paragraph(); pl.paragraph_format.keep_with_next = True; docx_link(pl, src_url(s))
        for f in FACTS.values():
            if f["source"] == sid and f["id"] in USED:
                d.add_paragraph("“" + f["quote"] + "”", style="List Bullet")
    cp = d.core_properties
    cp.title = R["title"]; cp.subject = R["subtitle"]; cp.author = "Redbelly DAO community research"
    cp.last_modified_by = "Redbelly DAO community research"; cp.comments = ""; cp.keywords = "Redbelly, Overledger, feasibility"
    from datetime import datetime, timezone
    cp.created = cp.modified = datetime(2026, 10, 15, tzinfo=timezone.utc); cp.revision = 1
    d.save(path)
    clean_app_props(path)

# ---------- Markdown ----------
def sources_md():
    L = ["# Source list", "", "Every factual claim in the report cites one of these sources as S1, S2 and so on. Under each source are the exact words the report relies on. "
         "Text snapshots of every page, with capture time and SHA-256, are in `evidence/` (see `evidence/manifest.json`).", ""]
    for sid in ORDER:
        s = SOURCES[sid]
        L.append(f"**S{NUM[sid]}. {s['publisher']}**: {s['title']}  ")
        L.append(src_url(s) + (f" (also `{s['url']}` in the repository)" if s.get("local") else ""))
        L.append("")
        for f in FACTS.values():
            if f["source"] == sid and f["id"] in USED:
                L.append(f"- \"{f['quote']}\"")
        L.append("")
    return "\n".join(L)

def summary_md():
    cs = R["community_summary"]
    L = ["# Community summary", "", "Plain language summary of the report, ready to share.", ""] + [p.replace("{site}", SITE) + "\n" for p in cs["paragraphs"]]
    L += ["## As an X thread", ""] + [t.replace("{site}", SITE) + "\n" for t in cs["thread"]]
    return "\n".join(L)

def write_evidence_index():
    """public/evidence/index.html, so the address /evidence opens a page instead of a 404."""
    man = json.loads((ROOT / "evidence/manifest.json").read_text())
    def row(m):
        orig = f'<a href="{esc(m["url"])}">{esc(m["url"])}</a>' if m.get("url") else "saved copy"
        return (f'<tr><td><a href="/{esc(m["file"])}">{esc(m["file"].split("/")[-1])}</a></td>'
                f'<td>{orig}</td><td>{esc(m.get("captured_utc") or "")}</td><td><code>{esc(m.get("sha256") or "")}</code></td></tr>')
    rows = "".join(row(m) for m in man)
    tests = "".join(f'<li><a href="/evidence/tests/{esc(f.name)}">{esc(f.name)}</a></li>' for f in sorted((ROOT / "evidence/tests").glob("*.txt")))
    page = (
        '<!doctype html><html lang="en-AU"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        '<title>Evidence | Could Redbelly connect to Quant Overledger?</title>'
        '<style>:root{color-scheme:light dark;--bg:#fff;--fg:#1c2a33;--mut:#5b6b76;--ln:#d7dee3;--acc:#B3272A}'
        '@media (prefers-color-scheme:dark){:root{--bg:#0c1317;--fg:#e8eef2;--mut:#9fb0bb;--ln:#25343d;--acc:#f0605f}}'
        'body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.55 system-ui,sans-serif}'
        'main{max-width:1100px;margin:0 auto;padding:32px 20px}a{color:var(--acc)}'
        '.w{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:13px}'
        'td,th{border-bottom:1px solid var(--ln);padding:8px 10px;text-align:left;vertical-align:top}'
        'code{font-size:11px;word-break:break-all;color:var(--mut)}</style></head><body><main>'
        '<p><a href="/">Back to the report</a> &middot; <a href="/sources">Source list</a></p>'
        '<h1>Evidence</h1>'
        f'<p>A saved text copy of every source page as it read on {esc(VERIFIED)}, the SHA-256 of each, and the logs of our own tests. '
        'The full list with hashes is in <a href="/evidence/manifest.json">manifest.json</a>.</p>'
        f'<h2>Our own test logs</h2><ul>{tests}</ul>'
        f'<h2>Source snapshots</h2><div class="w"><table><thead><tr><th>Snapshot</th><th>Original page</th><th>Captured (UTC)</th><th>SHA-256</th></tr></thead><tbody>{rows}</tbody></table></div>'
        '</main></body></html>')
    (PUB / "evidence" / "index.html").write_text(page, encoding="utf-8")

def prepare_assets():
    shutil.copy(ROOT / "scripts/site.css", PUB / "site.css")
    shutil.copy(ROOT / "scripts/site.js", PUB / "site.js")
    (PUB / "evidence/tests").mkdir(parents=True, exist_ok=True)
    for f in (ROOT / "evidence/tests").glob("*.txt"):
        shutil.copy(f, PUB / "evidence/tests" / f.name)
    # the saved copy of every source and the manifest of hashes are published with the site, so a reader can check any quote
    shutil.rmtree(PUB / "evidence/snapshots", ignore_errors=True)
    shutil.copytree(ROOT / "evidence/snapshots", PUB / "evidence/snapshots")
    shutil.copy(ROOT / "evidence/manifest.json", PUB / "evidence/manifest.json")
    write_evidence_index()
    from PIL import Image
    icon = PUB / "brand/icon.png"
    for size, name in ((32, "icon-32.png"), (180, "icon-180.png")):
        Image.open(icon).convert("RGBA").resize((size, size), Image.LANCZOS).save(PUB / "brand" / name, optimize=True)
    logo = PUB / "brand/dao-logo.png"
    im = Image.open(logo)
    if im.height > 200:
        im = im.convert("RGBA"); w = round(im.width * 176 / im.height)
        im.resize((w, 176), Image.LANCZOS).save(logo, optimize=True)

def make_og_image():
    """The 1200 x 630 link preview image (brand kit: the Open Graph size for any linked page)."""
    from playwright.sync_api import sync_playwright
    fonts = (PUB / "fonts").as_uri(); logo = (PUB / "brand/dao-logo.png").as_uri()
    page = (f'<!doctype html><html><head><meta charset="utf-8"><style>'
            f'@font-face{{font-family:"BV";src:url({fonts}/be-vietnam-pro-latin-700-normal.woff2);font-weight:700}}'
            f'@font-face{{font-family:"BV";src:url({fonts}/be-vietnam-pro-latin-400-normal.woff2);font-weight:400}}'
            f'@font-face{{font-family:"JM";src:url({fonts}/jetbrains-mono-latin-500-normal.woff2);font-weight:500}}'
            'html,body{margin:0}body{width:1200px;height:630px;background:#0c1317;color:#f1f5f8;font-family:BV,sans-serif;position:relative;overflow:hidden}'
            '.band{position:absolute;inset:40px;background:#1d3140;border-radius:8px;padding:56px 64px}'
            '.k{display:flex;align-items:center;gap:14px;font:500 22px JM,monospace;letter-spacing:.08em;text-transform:uppercase;color:#8fa1ad}'
            '.k i{display:block;width:40px;height:6px;border-radius:3px;background:#EF5350}'
            'h1{font-size:68px;line-height:1.06;letter-spacing:-.02em;margin:26px 0 22px;max-width:860px}'
            'p{font-size:28px;line-height:1.4;color:#c2ced6;margin:0;max-width:820px}'
            '.v{position:absolute;left:64px;bottom:52px;font:500 22px JM,monospace;color:#ffb3ae}'
            'img{position:absolute;right:64px;bottom:44px;height:84px}</style></head><body><div class="band">'
            '<div class="k"><i></i>Redbelly DAO community research</div>'
            f'<h1>{esc(R["title"])}</h1><p>Feasibility research from public sources. No partnership or talks between Redbelly and Quant have been announced.</p>'
            f'<div class="v">Sources last verified {VERIFIED}</div><img src="{logo}" alt=""></div></body></html>')
    tmp = ROOT / "build/og.html"; tmp.write_text(page)
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1200, "height": 630})
        pg.goto(tmp.as_uri()); pg.evaluate("document.fonts.ready.then(() => true)"); pg.wait_for_timeout(300)
        pg.screenshot(path=str(PUB / "brand/og.png")); b.close()

def main():
    (PUB / "downloads").mkdir(parents=True, exist_ok=True); (ROOT / "build").mkdir(exist_ok=True)
    prepare_assets()
    make_og_image()
    bodies = chapter_bodies()
    (PUB / "index.html").write_text(home_page(bodies))
    chap_dir = PUB / "chapters"
    shutil.rmtree(chap_dir, ignore_errors=True); chap_dir.mkdir()
    for cid in CH_IDS:
        (chap_dir / f"{CH_SLUG[cid]}.html").write_text(chapter_page(cid, bodies[cid]))
    (PUB / "sources.html").write_text(sources_page())
    (PUB / "404.html").write_text(not_found_page())
    (ROOT / "build/print.html").write_text(print_page())
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page()
        pg.goto((ROOT / "build/print.html").as_uri()); pg.wait_for_timeout(600)
        pg.pdf(path=str(PUB / "downloads" / PDF_NAME), format="A4", print_background=True, outline=True, tagged=True,
               margin={"top": "14mm", "bottom": "16mm", "left": "14mm", "right": "14mm"}, display_header_footer=True,
               header_template="<span></span>",
               footer_template="<div style='font-family:sans-serif;font-size:9px;width:100%;padding:0 14mm;display:flex;justify-content:space-between;color:#5b6b76'>"
                               "<span>Could Redbelly connect to Quant Overledger? Feasibility research only.</span><span><span class='pageNumber'></span> / <span class='totalPages'></span></span></div>")
        b.close()
    build_docx(PUB / "downloads" / DOCX_NAME, render_pngs())
    (ROOT / "SOURCES.md").write_text(sources_md()); shutil.copy(ROOT / "SOURCES.md", PUB / "downloads/SOURCES.md")
    (ROOT / "COMMUNITY-SUMMARY.md").write_text(summary_md()); shutil.copy(ROOT / "COMMUNITY-SUMMARY.md", PUB / "downloads/COMMUNITY-SUMMARY.md")
    print(f"built: {len(ORDER)} sources cited, {len(USED)} facts used, {len(FIGNUM)} figures")

if __name__ == "__main__":
    main()
