#!/usr/bin/env python3
"""Quality gate for the TASK-29 report. Runs every check below, prints a table and ends with RESULT: PASS or RESULT: FAIL.

  python3 scripts/check.py              offline checks
  python3 scripts/check.py --online     also asks every source link (and archived copy) whether it answers today
  python3 scripts/check.py --report     also writes the table to QUALITY-GATE.md, so a reviewer can see the result on GitHub

 1  every fact quote is found word for word in its saved snapshot, and every snapshot still matches its SHA-256 in
    evidence/manifest.json
 2  every citation points to a known fact, every source is cited, and every paragraph with a status label cites a fact
 3  no wording that suggests a partnership, talks, plans or an agreement between Redbelly and Quant
 4  no price talk, investment advice or hype words, no en dash, em dash or emoji
 5  no personal name, handle or personal website from the local, git-ignored .credit-names and .credit-links files in
    any deliverable, no "made by" line, and the Word file names no tool or person in its properties
 6  each X post fits in 280 characters (a link counts as 23)
 7  the PDF, the Word file and the website (home page and every chapter page) carry the current text of every paragraph,
    and the source page carries every quote
 8  every diagram follows the SVG rules: brand palette, prefixed ids, title and desc, no styles, images, scripts or
    outside links, and only words from its approved wording in content/diagrams.json (this catches typos)
 9  the website loads nothing from another domain and runs no inline script, so the Content-Security-Policy in
    vercel.json never blocks part of a page
10  every link inside the website opens a page that exists, at a clean address (no .html), and every #anchor exists
11  the chain IDs agree: what our RPC test answered, what Redbelly's documentation says, and every chain ID the report states
12  each test log was run on the date the report gives for it
"""
import argparse, functools, hashlib, html, json, re, subprocess, sys, urllib.request, zipfile
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "public"
PDF = PUB / "downloads/redbelly-overledger-feasibility-report.pdf"
DOCX = PUB / "downloads/redbelly-overledger-feasibility-report.docx"
CITE = re.compile(r"\{([a-z0-9 -]+)\}")
LABELLED = ("confirmed", "tested", "analysis", "unknown")
NOT_DELIVERED = {".git", "build", "__pycache__", ".venv", "venv", ".vercel", "node_modules"}

def load(rel):
    return json.loads((ROOT / rel).read_text())

R = load("content/report.json")
FACT_LIST = load("content/facts.json")
F = {f["id"]: f for f in FACT_LIST}
SOURCES = {s["id"]: s for s in load("content/sources.json")}
DIAGRAMS = load("content/diagrams.json")
MANIFEST = {m["file"]: m for m in load("evidence/manifest.json")}
NOTES = []

def norm(s):
    """Curly quotes and odd spaces made plain, runs of white space made one space."""
    for fancy, simple in (("’", "'"), ("‘", "'"), ("“", '"'), ("”", '"'), ("\u00a0", " "), ("\u202f", " ")):
        s = s.replace(fancy, simple)
    return re.sub(r"\s+", " ", s).strip()

def squash(t):
    """Letters and digits only, for comparing text across PDF, Word and HTML line breaks."""
    return re.sub(r"[^a-z0-9]", "", norm(t).lower())

def text_of_html(markup):
    return html.unescape(re.sub(r"<[^>]+>", " ", markup))

def snapshot(source):
    return ROOT / "evidence/snapshots" / source["file"]

# ---------- our own words, gathered once ----------
def gather_prose():
    """(where, text) for every piece of text we wrote, in reading order, and the community summary separately."""
    prose = [("status note", R["status_note"])]
    for k in R["key_findings"]:
        prose += [("finding", k["q"]), ("finding", k["a"] + " {" + k["cite"] + "}")]
    for s in R["sections"]:
        prose.append((s["id"] + " title", re.sub(r"^\d+\.\s*", "", s["title"])))
        if s.get("intro"):
            prose.append((s["id"] + " intro", s["intro"]))
        for i, b in enumerate(s["blocks"], 1):
            rows = [r[0] + " " + r[1] + " {" + r[2] + "}" for r in b.get("rows", [])]
            prose += [(f"{s['id']} block {i}", t) for t in [b.get("heading", ""), b.get("text", "")] + b.get("list", []) + rows if t]
    for term, d in R.get("glossary", []):
        prose += [("glossary " + term, term), ("glossary " + term, d)]
    for d in DIAGRAMS:
        prose += [(f"figure {d['id']} title", d["title"]), (f"figure {d['id']} caption", d["caption"])]
    cs = R["community_summary"]
    summary = [("community summary", p) for p in cs["paragraphs"]] + [("X thread", t) for t in cs["thread"]]
    return prose, summary

def diagram_files():
    """(diagram, variant, tag, path) for the wide and narrow SVG of every diagram."""
    return [(d, v, tag, PUB / "diagrams" / f"{d['id']}-{v}.svg") for d in DIAGRAMS for v, tag in (("wide", "w"), ("narrow", "n"))]

def drawn_words(svg, tags="text|title|desc"):
    return " ".join(re.sub(r"<[^>]+>", " ", m) for m in re.findall(rf"<(?:{tags})\b.*?</(?:{tags})>", svg, re.S))

PROSE, SUMMARY = gather_prose()
SVG_TEXTS = [(f"figure {d['id']} {v} text", html.unescape(re.sub(r"\s+", " ", drawn_words(p.read_text()))))
             for d, v, _, p in diagram_files() if p.exists()]
USED = {fid for _, t in PROSE for m in CITE.finditer(t) for fid in m.group(1).split()}
# our words without the citation markers: what the wording rules (3, 4, 5, 11) read
OURS = [(w, CITE.sub("", t)) for w, t in PROSE + SUMMARY] + SVG_TEXTS

# ---------- the checks ----------
CHECKS = []

def check(title):
    def register(fn):
        CHECKS.append((title, fn))
        return fn
    return register

@check("Quotes are word for word in their saved snapshots")
def quotes_and_snapshots():
    errors = []
    for rel, m in MANIFEST.items():
        p = ROOT / rel
        if not p.exists():
            errors.append(f"evidence: {rel} listed in manifest but missing")
        elif hashlib.sha256(p.read_bytes()).hexdigest() != m["sha256"]:
            errors.append(f"evidence: {rel} changed after capture (SHA-256 mismatch)")
    for s in SOURCES.values():
        rel = str(snapshot(s).resolve().relative_to(ROOT))
        if rel not in MANIFEST:
            errors.append(f"source {s['id']}: snapshot {rel} not in manifest")
    for f in FACT_LIST:
        s = SOURCES.get(f["source"])
        if not s:
            errors.append(f"fact {f['id']}: unknown source {f['source']}")
            continue
        # Markdown bold markers in a copied README are formatting, not words
        page = snapshot(s).read_text(encoding="utf-8", errors="ignore").replace("**", "")
        if norm(f["quote"]) not in norm(page):
            errors.append(f"fact {f['id']}: quote not found in {s['file']}")
    return errors, f"{len(FACT_LIST)} quotes, {len(MANIFEST)} evidence files"

@check("Every claim cites a known fact and every source is cited")
def citations_and_labels():
    errors = []
    for s in R["sections"]:
        for i, b in enumerate(s["blocks"], 1):
            where = f"{s['id']} block {i}"
            if b["label"] not in LABELLED + ("summary_table", "method"):
                errors.append(f"{where}: unknown label {b['label']}")
            if b["label"] in LABELLED:
                for t in ([b["text"]] if b.get("text") else []) + b.get("list", []):
                    if not CITE.search(t):
                        errors.append(f"{where}: a {b['label']} statement cites no source: {t[:60]}")
    for d in DIAGRAMS:
        if not CITE.search(d["caption"]):
            errors.append(f"figure {d['id']}: caption cites no source")
    for where, t in PROSE:
        for m in CITE.finditer(t):
            errors += [f"{where}: cites unknown fact {fid}" for fid in m.group(1).split() if fid not in F]
    NOTES.extend(f"fact {fid} is not cited (kept in the fact list only)" for fid in F if fid not in USED)
    cited_sources = {F[fid]["source"] for fid in USED if fid in F}
    errors += [f"source {sid} is never cited" for sid in SOURCES if sid not in cited_sources]
    if R["key_findings"] and not all(k["label"] in LABELLED for k in R["key_findings"]):
        errors.append("key findings: every row needs a status label")
    return errors, f"{len(USED)} facts cited from {len(cited_sources)} sources"

@check("No wording that suggests a Redbelly and Quant link")
def partnership_wording():
    errors = []
    link_words = re.compile(r"partner|talks\b|negotiat|collaborat|agreement|in discussion|working with|plans?\b|roadmap|"
                            r"will (integrate|connect|join|support|list)", re.I)
    negated = re.compile(r"\b(no|not|none|nothing|never|without)\b", re.I)
    for where, t in OURS:
        for sentence in re.split(r"(?<=[.?!])\s+", t):
            if sentence.rstrip().endswith("?"):
                continue   # a question asks, it does not claim
            if link_words.search(sentence) and re.search(r"redbelly|quant", sentence, re.I) and not negated.search(sentence) \
               and "The Clearing House" not in sentence:
                errors.append(f"{where}: wording could suggest a Redbelly and Quant link: {sentence[:90]}")
    return errors, f"{len(OURS)} passages"

HYPE = ("price", "prices", "priced", "profit", "profits", "returns?", "yield", "apy", "apr", "buy", "sell", "bullish", "bearish",
        "moon", "pump", "undervalued", "investment advice", "financial advice", "guaranteed", "revolutionary", "game.?changer",
        "unleash", "seamless", "skyrocket", "massive")
EMOJI = re.compile("[\U0001F300-\U0001FAFF\u2600-\u27bf\u2b00-\u2bff\ufe0f]")

@check("No price talk, hype words, dashes or emoji")
def tone_and_typography():
    errors = []
    banned = re.compile(r"\b(" + "|".join(HYPE) + r")\b", re.I)
    for where, t in OURS:
        m = banned.search(t)
        if m:
            errors.append(f"{where}: banned word '{m.group(0)}'")
    for where, t in OURS + [("fact " + f["id"], f["quote"]) for f in FACT_LIST]:
        if "\u2013" in t or "\u2014" in t:
            errors.append(f"{where}: contains an en dash or em dash")
        if EMOJI.search(t):
            errors.append(f"{where}: contains an emoji or symbol")
    return errors, f"{len(OURS)} passages, {len(FACT_LIST)} quotes"

def private_list(name):
    p = ROOT / name
    return [line.strip().lower() for line in p.read_text().splitlines() if line.strip()] if p.exists() else []

def readable_text(p):
    if p.suffix == ".pdf":
        return subprocess.run(["pdftotext", str(p), "-"], capture_output=True, text=True).stdout
    if p.suffix == ".docx":
        from docx import Document
        return "\n".join(x.text for x in Document(p).paragraphs)
    return p.read_bytes().decode("utf-8", "ignore")

@check("No personal credit in any deliverable")
def personal_credit():
    errors = []
    names = private_list(".credit-names") + private_list(".credit-links") if (ROOT / ".credit-names").exists() else []
    if (ROOT / ".credit-names").exists():
        files = [p for p in ROOT.rglob("*") if p.is_file() and not NOT_DELIVERED & set(p.relative_to(ROOT).parts)
                 and p.name not in (".credit-names", ".credit-links") and p.suffix != ".woff2"]
        for p in files:
            body = readable_text(p).lower()
            if any(n in body for n in names):
                errors.append(f"personal credit: {p.relative_to(ROOT)} contains a listed name or link")
        if re.search(r"made by|created by|author:", " ".join(t for _, t in OURS), re.I):
            errors.append("personal credit: a 'made by' style line")
        detail = f"{len(files)} files scanned for {len(names)} private entries"
    else:
        NOTES.append("personal credit scan skipped, .credit-names not present")
        detail = "name scan skipped (no .credit-names here), Word properties checked"
    if DOCX.exists():
        from docx import Document
        props = Document(DOCX).core_properties
        if re.search(r"<Application>[^<]+</Application>", zipfile.ZipFile(DOCX).read("docProps/app.xml").decode("utf-8", "ignore")):
            errors.append("Word file: docProps/app.xml names an application")
        for field in ("author", "last_modified_by", "comments"):
            value = (getattr(props, field) or "").lower()
            if "python-docx" in value or any(n in value for n in names):
                errors.append(f"Word file: the {field} property names a tool or a person")
    return errors, detail

@check("Every X post fits in 280 characters")
def x_thread_length():
    thread = R["community_summary"]["thread"]
    lengths = [len(t.replace("{site}", "x" * 23)) for t in thread]
    errors = [f"X post {i} is {n} characters, over 280" for i, n in enumerate(lengths, 1) if n > 280]
    return errors, f"{len(thread)} posts, longest {max(lengths, default=0)}"

@check("PDF, Word file and website carry the current text")
def outputs_up_to_date():
    errors = []
    outputs = {}
    if PDF.exists():
        outputs["PDF"] = squash(subprocess.run(["pdftotext", "-raw", str(PDF), "-"], capture_output=True, text=True).stdout)
    if DOCX.exists():
        from docx import Document
        d = Document(DOCX)
        outputs["Word file"] = squash("\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]))
    home, chapters = PUB / "index.html", sorted((PUB / "chapters").glob("*.html"))
    if home.exists() and chapters:
        outputs["web page"] = squash(text_of_html("\n".join(p.read_text() for p in [home] + chapters)))
    for name in ("PDF", "Word file", "web page"):
        if name not in outputs:
            errors.append(f"{name} missing, run scripts/build.py")
            continue
        for where, t in PROSE + SUMMARY:
            # the words before a paragraph's first citation stand for the paragraph
            opening = CITE.split(t)[0]
            probe = squash(opening.replace("{site}", ""))[:60]
            if len(probe) >= 12 and probe not in outputs[name]:
                errors.append(f"{name} is out of date: '{opening[:50]}' ({where}) not found, run scripts/build.py")
                break
    source_page = PUB / "sources.html"
    if source_page.exists():
        listed = norm(text_of_html(source_page.read_text()))
        missing = [fid for fid in USED if fid in F and norm(F[fid]["quote"]) not in listed]
        if missing:
            errors.append(f"source page is out of date: quote {missing[0]} missing, run scripts/build.py")
    else:
        errors.append("public/sources.html missing, run scripts/build.py")
    return errors, f"{len(PROSE + SUMMARY)} paragraphs in {len(outputs)} outputs"

PALETTE = {"#0c1317", "#141f26", "#17222a", "#1d3140", "#263641", "#223039", "#f1f5f8", "#c2ced6", "#8fa1ad", "#ffb3ae", "#ef5350",
           "#86efac", "#fcd34d", "#ffffff", "#fff"}

@check("Diagrams follow the SVG rules and their approved wording")
def diagram_rules():
    errors = []
    files = diagram_files()
    for d, v, tag, p in files:
        if not p.exists():
            errors.append(f"figure {d['id']}: {p.name} missing")
            continue
        svg = p.read_text()
        errors += [f"figure {d['id']} {v}: contains {bad}" for bad in
                   ("<style", "<image", "<foreignObject", "<script", "class=", "xlink:href=\"http", "href=\"http", "@import") if bad in svg]
        prefix = f"{d['id']}-{tag}-"
        if f'id="{prefix}title"' not in svg or "<desc" not in svg:
            errors.append(f"figure {d['id']} {v}: needs <title id=\"{prefix}title\"> and <desc>")
        errors += [f"figure {d['id']} {v}: id '{i}' is not prefixed with {prefix}" for i in re.findall(r'\sid="([^"]+)"', svg) if not i.startswith(prefix)]
        errors += [f"figure {d['id']} {v}: colour {c} is outside the brand palette" for c in re.findall(r"#[0-9a-fA-F]{3,6}\b", svg) if c.lower() not in PALETTE]
        # every word drawn in the figure must come from its approved wording (this catches typos)
        allowed = set(re.findall(r"[a-z0-9]+", html.unescape(" ".join(d["content"]) + " " + d["title"] + " " + d["caption"]).lower()))
        drawn = set(re.findall(r"[a-z0-9]+", html.unescape(drawn_words(svg, "text")).lower()))
        errors += [f"figure {d['id']} {v}: the word '{w}' is drawn but not in the figure's wording in content/diagrams.json"
                   for w in sorted(drawn - allowed)]
    return errors, f"{len(DIAGRAMS)} diagrams, {len(files)} SVG files"

@check("Website loads nothing from outside and runs no inline script")
def self_contained_site():
    errors = []
    files = sorted(PUB.rglob("*.html")) + [PUB / "site.css", PUB / "site.js"]
    for p in (f for f in files if f.exists()):
        t = p.read_text()
        for m in re.finditer(r"<(script|link|img|iframe|source|video|audio)\b[^>]*\b(src|href)=[\"'](https?:)?//", t, re.I):
            if 'rel="canonical"' not in m.group(0):   # a canonical link names the page's own address and loads nothing
                errors.append(f"{p.name}: loads an outside resource ({m.group(0)[:60]})")
        if re.search(r"url\(\s*[\"']?(https?:)?//", t):
            errors.append(f"{p.name}: CSS loads an outside resource")
        if "\u2013" in t or "\u2014" in t:
            errors.append(f"{p.name}: contains an en dash or em dash")
        if p.suffix == ".html":
            # script-src 'self' allows only script files from this site; a JSON data block is not run, so it is fine
            for m in re.finditer(r"<script\b([^>]*)>", t, re.I):
                attrs = m.group(1)
                if "src=" not in attrs and 'type="application/json"' not in attrs:
                    errors.append(f"{p.relative_to(ROOT)}: inline script, which the Content-Security-Policy would block")
            if re.search(r"\son[a-z]+\s*=\s*[\"']|href=[\"']javascript:", t, re.I):
                errors.append(f"{p.relative_to(ROOT)}: inline event handler or javascript: link, which the Content-Security-Policy would block")
    csp = [h["value"] for rule in load("vercel.json").get("headers", []) for h in rule["headers"] if h["key"] == "Content-Security-Policy"]
    if not csp or "script-src 'self'" not in csp[0]:
        errors.append("vercel.json: no Content-Security-Policy with script-src 'self'")
    return errors, f"{len(files)} files"

@functools.lru_cache(maxsize=None)
def page_ids(p):
    return set(re.findall(r'\sid="([^"]+)"', p.read_text()))

def page_for(path):
    """The file Vercel serves for a site path (cleanUrls: /sources serves sources.html), or None."""
    path = path.split("?")[0]
    if path == "/":
        return PUB / "index.html"
    f = PUB / path.lstrip("/")
    return next((c for c in (f, PUB / (path.lstrip("/") + ".html"), f / "index.html") if c.is_file()), None)

@check("Links inside the website resolve at clean addresses")
def internal_links():
    errors = []
    pages = sorted(PUB.rglob("*.html"))
    if not (PUB / "404.html").exists():
        errors.append("public/404.html missing, run scripts/build.py")
    count = 0
    for p in pages:
        rel = p.relative_to(ROOT)
        for m in re.finditer(r'\s(href|src)="([^"]+)"', p.read_text()):
            u = html.unescape(m.group(2))
            if u.startswith(("https://", "http://", "mailto:")):
                continue
            count += 1
            if u.startswith("#"):
                if len(u) > 1 and u[1:] not in page_ids(p):
                    errors.append(f"{rel}: link {u} points to an anchor that is not on the page")
                continue
            if not u.startswith("/"):
                errors.append(f"{rel}: link {u} is relative, use a path from the site root")
                continue
            path, _, anchor = u.partition("#")
            if path.split("?")[0].endswith(".html"):
                errors.append(f"{rel}: link {u} ends in .html, use the clean address")
            target = page_for(path)
            if target is None:
                errors.append(f"{rel}: link {u} opens nothing")
            elif anchor and target.suffix == ".html" and anchor not in page_ids(target):
                errors.append(f"{rel}: link {u} points to an anchor that is not on {target.name}")
    return errors, f"{count} links on {len(pages)} pages"

def test_log(source_id):
    return (ROOT / "evidence/snapshots" / SOURCES[source_id]["file"]).resolve()

@check("Chain IDs agree: our RPC test, Redbelly's docs and the report")
def chain_ids():
    errors = []
    documented = {}   # network -> chain ID quoted from Redbelly's documentation
    for f in FACT_LIST:
        m = re.search(r"governors\.(mainnet|testnet)\.redbelly\.network\s+Chain ID:\s*(\d+)", f["quote"])
        if m:
            documented[m.group(1)] = int(m.group(2))
    if set(documented) != {"mainnet", "testnet"}:
        return [f"facts: need the documented chain ID of mainnet and testnet, found {sorted(documented) or 'none'}"], ""
    answers = 0
    for line in test_log("test-rpc").read_text().splitlines():
        m = re.match(r"(https://\S+) eth_chainId -> (\{.*\})\s*$", line)
        if not m:
            continue
        answers += 1
        network = "testnet" if "testnet" in m.group(1) else "mainnet"
        got = int(json.loads(m.group(2))["result"], 16)
        if got != documented[network]:
            errors.append(f"test log: {m.group(1)} answered chain ID {got}, Redbelly documents {documented[network]} for {network}")
    if not answers:
        errors.append("test log: no eth_chainId answer found in the RPC log")
    stated = 0
    for where, t in OURS:
        for m in re.finditer(r"chain ID (\d+)(?: \((0x[0-9a-f]+)\))?", t, re.I):
            stated += 1
            n = int(m.group(1))
            if n not in documented.values():
                errors.append(f"{where}: states chain ID {n}, which is neither mainnet ({documented['mainnet']}) nor testnet ({documented['testnet']})")
            if m.group(2) and int(m.group(2), 16) != n:
                errors.append(f"{where}: chain ID {n} is not {m.group(2)}")
    return errors, f"{answers} RPC answers, {stated} statements in the report"

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

def long_date(iso):
    d = datetime.strptime(iso[:10], "%Y-%m-%d")
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"

@check("Each test log was run on the date the report gives")
def test_dates():
    errors = []
    logs = [s for s in SOURCES.values() if s.get("local")]
    for s in logs:
        first = re.search(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", test_log(s["id"]).read_text())
        if not first:
            errors.append(f"source {s['id']}: the log carries no run time")
            continue
        ran = long_date(first.group(0))
        if ran not in s["title"]:
            errors.append(f"source {s['id']}: the log ran on {ran}, the title says otherwise: {s['title']}")
        if ran not in R["labels"]["tested"]:
            errors.append(f"labels: the Tested label does not give {ran}, the day {s['id']} ran")
    return errors, f"{len(logs)} test logs"

def online_links():
    """Notes only: a page that refuses a script can still be fine in a browser."""
    for s in SOURCES.values():
        if s.get("local"):
            continue
        for label, url in [("link", s["url"])] + ([("archived copy", s["archived"])] if s.get("archived") else []):
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                code = urllib.request.urlopen(req, timeout=30).status
            except Exception as e:
                code = getattr(e, "code", str(e)[:40])
            NOTES.append(f"{label} {s['id']}: 200" if code == 200 else f"{label} {s['id']}: answered {code} (check in a browser: {url})")

# ---------- run ----------
def write_report(rows, passed):
    when = datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")
    lines = ["# Quality gate", "",
             f"Result of `python3 scripts/check.py` on {when}: **{'PASS' if passed else 'FAIL'}**, "
             f"{sum(1 for r in rows if r[2] == 'PASS')} of {len(rows)} checks passed.", "",
             "Each check is described at the top of `scripts/check.py`. GitHub Actions runs the same script on every push "
             "(`.github/workflows/check.yml`), without the private name lists, so that run skips the name scan in check 5.", "",
             "| # | Check | Result | What was checked |", "|---|---|---|---|"]
    lines += [f"| {n} | {title} | {result} | {detail} |" for n, title, result, detail in rows]
    if NOTES:
        lines += ["", "Notes:", ""] + [f"- {n}" for n in NOTES]
    (ROOT / "QUALITY-GATE.md").write_text("\n".join(lines) + "\n")

def main():
    ap = argparse.ArgumentParser(description="Quality gate for the report.")
    ap.add_argument("--online", action="store_true", help="also ask every source link whether it answers today")
    ap.add_argument("--report", action="store_true", help="also write the results table to QUALITY-GATE.md")
    args = ap.parse_args()
    rows, failures = [], []
    for n, (title, fn) in enumerate(CHECKS, 1):
        errors, detail = fn()
        failures += errors
        rows.append((n, title, "PASS" if not errors else f"FAIL ({len(errors)})", detail))
    if args.online:
        online_links()
    for n in NOTES:
        print("note:", n)
    for e in failures:
        print("FAIL:", e)
    print()
    for n, title, result, detail in rows:
        print(f"{n:>2}  {title:<64} {result:<9} {detail}")
    print()
    print(f"{len(FACT_LIST)} facts, {len(SOURCES)} sources, {len(MANIFEST)} evidence files checked")
    if args.report:
        write_report(rows, not failures)
        print("wrote QUALITY-GATE.md")
    print("RESULT:", "PASS" if not failures else f"FAIL ({len(failures)})")
    sys.exit(1 if failures else 0)

if __name__ == "__main__":
    main()
