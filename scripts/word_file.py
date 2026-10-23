#!/usr/bin/env python3
"""The Word version of the report, built with python-docx. The diagrams go in as PNGs rendered from their SVG with Playwright."""
import re, zipfile
from report_model import *
from pages import svg_inline

def render_pngs():
    out = {}
    from playwright.sync_api import sync_playwright
    fonts = (PUB / "fonts").as_uri()
    face = "".join(f'@font-face{{font-family:"{fam}";src:url({fonts}/{f});font-weight:{w}}}' for fam, f, w in
                   [("Be Vietnam Pro", "be-vietnam-pro-latin-400-normal.woff2", 400), ("Be Vietnam Pro", "be-vietnam-pro-latin-500-normal.woff2", 500),
                    ("Be Vietnam Pro", "be-vietnam-pro-latin-600-normal.woff2", 600), ("Be Vietnam Pro", "be-vietnam-pro-latin-700-normal.woff2", 700),
                    ("JetBrains Mono", "jetbrains-mono-latin-400-normal.woff2", 400), ("JetBrains Mono", "jetbrains-mono-latin-500-normal.woff2", 500)])
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1000, "height": 800}, device_scale_factor=2)
        for did in FIGNUM:
            tmp = ROOT / "build" / f"fig-{did}.html"
            tmp.write_text(f'<html><head><meta charset="utf-8"><style>{face}body{{margin:0;background:#141f26}}.w{{width:960px;padding:20px}}svg{{display:block;width:100%;height:auto}}</style></head>'
                           f'<body><div class="w">{svg_inline(did, "wide")}</div></body></html>')
            pg.goto(tmp.as_uri()); pg.evaluate("document.fonts.ready.then(() => true)"); pg.wait_for_timeout(300)
            f = ROOT / "build" / f"fig-{did}.png"; pg.locator(".w").screenshot(path=str(f)); out[did] = f
        b.close()
    return out

LINK_RED = "B3272A"
DOC_AUTHOR = "Redbelly DAO community research, TASK-29"

def xml_el(tag, attrs=None, children=(), text=None):
    """A WordprocessingML element, such as xml_el("w:u", {"w:val": "single"})."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement(tag)
    for key, value in (attrs or {}).items():
        el.set(qn(key), value)
    for child in children:
        el.append(child)
    if text is not None:
        el.text = text
    return el

def add_hyperlink(paragraph, url, text=None):
    """Appends a clickable link (red, underlined) to a Word paragraph; python-docx has no call for this."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    style = xml_el("w:rPr", children=[xml_el("w:color", {"w:val": LINK_RED}), xml_el("w:u", {"w:val": "single"})])
    run = xml_el("w:r", children=[style, xml_el("w:t", {"xml:space": "preserve"}, text=text or url)])
    paragraph._p.append(xml_el("w:hyperlink", {"r:id": rid}, [run]))

# python-docx's template carries a thumbnail and names a word processor and company; none of that belongs to this report
TEMPLATE_TRACES = {
    "docProps/app.xml": [(rb"<Application>.*?</Application>", b"<Application></Application>"), (rb"<AppVersion>.*?</AppVersion>", b""),
                         (rb"<Company>.*?</Company>", b"<Company></Company>")],
    "_rels/.rels": [(rb"<Relationship[^>]*thumbnail[^>]*/>", b"")],
    "[Content_Types].xml": [(rb"<Override[^>]*thumbnail[^>]*/>", b"")],
}

def strip_template_traces(path):
    with zipfile.ZipFile(path) as z:
        parts = [(info, z.read(info.filename)) for info in z.infolist() if info.filename != "docProps/thumbnail.jpeg"]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in parts:
            for pattern, replacement in TEMPLATE_TRACES.get(info.filename, []):
                data = re.sub(pattern, replacement, data)
            z.writestr(info, data)

def build_docx(path, pngs):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.4)
    for name in ("Title", "Heading 1", "Heading 2"):
        d.styles[name].font.color.rgb = RGBColor(0x0F, 0x1A, 0x22)
    st.font.size = Pt(11)
    d.add_heading(R["title"], 0)
    d.add_paragraph(R["subtitle"] + ".").runs[0].italic = True
    p0 = d.add_paragraph(f"Redbelly DAO community research. Feasibility report, sources last verified {VERIFIED}. Web version: "); add_hyperlink(p0, SITE)
    d.add_paragraph(plain(R["status_note"]))
    d.add_heading("How to read the labels", 1)
    for k, v in R["labels"].items():
        p = d.add_paragraph(style="List Bullet"); r = p.add_run(LABELS[k] + ": "); r.bold = True; p.add_run(v)
    d.add_heading("Plain words", 1)
    d.add_paragraph("The technical terms this report uses, in one or two lines each. Where a term is Overledger's or Redbelly's own, the source is cited.")
    for term, dfn in R["glossary"]:
        p = d.add_paragraph(style="List Bullet"); r = p.add_run(term + ": "); r.bold = True; p.add_run(plain(dfn))
    d.add_heading("Key findings", 1)
    t = d.add_table(rows=1, cols=4); t.style = "Light Grid Accent 2"
    for i, x in enumerate(("Question", "Short answer", "Detail", "Status")):
        t.rows[0].cells[i].text = x
    for k in R["key_findings"]:
        c = t.add_row().cells; c[0].text = k["q"]; c[1].text = k["verdict"]; c[2].text = plain(k["a"] + " {" + k["cite"] + "}"); c[3].text = LABELS[k["label"]]
    def add_fig(did):
        d.add_picture(str(pngs[did]), width=Cm(16))
        d.inline_shapes[-1]._inline.docPr.set("descr", f"Figure {FIGNUM[did]}. {DIAGRAMS[did]['title']}. " + plain(DIAGRAMS[did]["caption"]))
        cap = d.add_paragraph(); r = cap.add_run(f"Figure {FIGNUM[did]}. {DIAGRAMS[did]['title']}. "); r.bold = True
        cap.add_run(plain(DIAGRAMS[did]["caption"]))
    for s in R["sections"]:
        d.add_heading(f"{SEC_NUM[s['id']]}. {strip_num(s['title'])}", 1)
        if s.get("intro"):
            d.add_paragraph(plain(s["intro"]))
        for did in figs_for(s["id"], -1):
            add_fig(did)
        for bi, b in enumerate(s["blocks"]):
            if b["label"] == "summary_table":
                t = d.add_table(rows=1, cols=2); t.style = "Light Grid Accent 2"
                t.rows[0].cells[0].text = "Question"; t.rows[0].cells[1].text = "Answer today"
                for r in b["rows"]:
                    c = t.add_row().cells; c[0].text = r[0]; c[1].text = plain(r[1] + " {" + r[2] + "}")
            else:
                if b.get("heading"):
                    d.add_heading(b["heading"], 2)
                tag = LABELS.get(b["label"], "Method")
                if b.get("text"):
                    p = d.add_paragraph(); r = p.add_run(f"[{tag}] "); r.bold = True; r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
                    p.add_run(plain(b["text"]))
                if b.get("list"):
                    if not b.get("text"):
                        r = d.add_paragraph().add_run(f"[{tag}]"); r.bold = True; r.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
                    # numbered lists carry their own numbers, so they match the figures (Word's auto numbering would run on across lists)
                    numbered = s["id"] == "compliance" and b.get("heading", "").startswith("Where")
                    for k, li in enumerate(b["list"], 1):
                        if numbered:
                            q = d.add_paragraph(f"{k}. " + plain(li)); q.paragraph_format.left_indent = Cm(0.6); q.paragraph_format.first_line_indent = Cm(-0.6)
                        else:
                            d.add_paragraph(plain(li), style="List Bullet")
            for did in figs_for(s["id"], bi):
                add_fig(did)
    cs = R["community_summary"]
    d.add_heading(f"{len(R['sections']) + 1}. Community summary", 1)
    for p in cs["paragraphs"]:
        d.add_paragraph(p.replace("{site}", SITE))
    d.add_heading("As an X thread", 2)
    for x in cs["thread"]:
        d.add_paragraph(x.replace("{site}", SITE))
    d.add_heading("Source list", 1)
    d.add_paragraph("Every factual claim cites one of these sources. Under each source are the exact words the report relies on.")
    pr = d.add_paragraph("Saved copies of every source page, with SHA-256 hashes: "); add_hyperlink(pr, f"{SITE}/evidence")
    for sid in ORDER:
        s = SOURCES[sid]
        p = d.add_paragraph(); r = p.add_run(f"S{NUM[sid]}  {s['publisher']}: {s['title']}"); r.bold = True
        p.paragraph_format.keep_with_next = True
        pl = d.add_paragraph(); pl.paragraph_format.keep_with_next = True; add_hyperlink(pl, src_url(s))
        if s.get("archived"):
            pa = d.add_paragraph("Archived copy: "); pa.paragraph_format.keep_with_next = True; add_hyperlink(pa, s["archived"])
        for f in FACTS.values():
            if f["source"] == sid and f["id"] in USED:
                d.add_paragraph("“" + f["quote"] + "”", style="List Bullet")
    cp = d.core_properties
    cp.title, cp.subject, cp.keywords, cp.comments = R["title"], R["subtitle"], "Redbelly, Overledger, feasibility", ""
    cp.author = cp.last_modified_by = DOC_AUTHOR
    from datetime import datetime, timezone
    cp.created = cp.modified = datetime(2026, 10, 15, tzinfo=timezone.utc); cp.revision = 1
    d.save(path)
    strip_template_traces(path)
